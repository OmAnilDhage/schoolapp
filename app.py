from flask import Flask, render_template, request, redirect, session
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = Flask(__name__)
app.secret_key = "school_secret_key"

FILE_NAME = "School_Progress_System.xlsx"
PASSWORD = "om@2025-26"
PASS_PERCENT = 35

def setup_excel():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Student_Record"
        sheet.append([
            "Date", "Class", "Roll No", "Name",
            "Exam Type", "English", "Hindi",
            "Marathi", "Maths", "Science",
            "Social Science", "Total",
            "Percentage", "Grade", "Status", "Remarks"
        ])
        wb.save(FILE_NAME)

def calculate_grade(p):
    if p >= 90: return "A+", "PASS"
    elif p >= 75: return "A", "PASS"
    elif p >= 60: return "B", "PASS"
    elif p >= 45: return "C", "PASS"
    elif p >= PASS_PERCENT: return "D", "PASS"
    else: return "F", "FAIL"

def calculate_remarks(p):
    if p >= 90: return "Excellent"
    elif p >= 75: return "Very Good"
    elif p >= 60: return "Good"
    elif p >= 35: return "Satisfactory"
    else: return "Needs Improvement"

@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if request.form["password"] == PASSWORD:
            session["logged_in"] = True
            return redirect("/form")
        return "Incorrect Password!"
    return render_template("login.html")

@app.route("/form", methods=["GET", "POST"])
def form():
    if not session.get("logged_in"):
        return redirect("/")
    
    if request.method == "POST":
        name = request.form["name"]
        s_class = request.form["class"]
        roll = request.form["roll"]
        exam_type = request.form["exam_type"]
        unit_total = int(request.form.get("unit_total", 20))

        if exam_type == "Term":
            sub_limit, part_limit, internal, max_total = 80, 40, 20, 600
        else:
            sub_limit = unit_total
            part_limit = sub_limit / 2
            internal, max_total = 0, 6 * sub_limit

        # Calculating Subject Totals
        eng = int(request.form["eng"]) + internal
        hin = int(request.form["hin"]) + internal
        mar = int(request.form["mar"]) + internal
        math = int(request.form["mat1"]) + int(request.form["mat2"]) + internal
        sci = int(request.form["sci1"]) + int(request.form["sci2"]) + internal
        soc = int(request.form["hist"]) + int(request.form["geo"]) + internal

        total = eng + hin + mar + math + sci + soc
        percent = (total / max_total) * 100
        grade, status = calculate_grade(percent)
        remarks = calculate_remarks(percent)

        setup_excel()
        wb = load_workbook(FILE_NAME)
        sheet = wb["Student_Record"]
        result_data = [
            datetime.now().strftime("%d-%m-%Y"), s_class, roll, name,
            exam_type, eng, hin, mar, math, sci, soc, total,
            round(percent, 2), grade, status, remarks
        ]
        sheet.append(result_data)
        wb.save(FILE_NAME)
        
        return render_template("result.html", s=result_data)
    
    return render_template("form.html")

if __name__ == "__main__":
    setup_excel()
    app.run(debug=True, port=5000)